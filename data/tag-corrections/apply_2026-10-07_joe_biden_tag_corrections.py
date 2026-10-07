#!/usr/bin/env python3
"""Validate and narrowly apply the user-reviewed Joe Biden historical tag corrections."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import sys
from pathlib import Path

from docx import Document

PROJECT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT))

from safe_batch import load_ledger, write_ledger, write_project_from_ledger  # noqa: E402

LEDGER = PROJECT / "data" / "article-ledger.json"
UPLOADED_DOCX = Path("/home/ubuntu/upload/ctcrazies_Editiedjoe_biden_tag_article_list_2026-10-07.docx")
ARCHIVED_DOCX = PROJECT / "data" / "tag-corrections" / "ctcrazies_Editiedjoe_biden_tag_article_list_2026-10-07.docx"
PLAN = PROJECT / "data" / "tag-corrections" / "2026-10-07-joe-biden-historical-tag-corrections.json"
REPORT = PROJECT / "data" / "tag-corrections" / "2026-10-07-joe-biden-historical-tag-corrections.md"
EXPECTED_HEADERS = ["NUM", "Exact X-Post Headline", "Current Tags"]


def parse_tags(cell_text: str, num: str) -> list[str]:
    tags = [value.strip() for value in cell_text.replace("\n", " ").split(",") if value.strip()]
    assert tags, f"NUM {num}: Current Tags cell is empty"
    assert len(tags) == len(set(tags)), f"NUM {num}: duplicate tags in returned DOCX: {tags}"
    assert 2 <= len(tags) <= 4, f"NUM {num}: expected 2–4 final tags, found {len(tags)}"
    return tags


def validate() -> dict:
    assert UPLOADED_DOCX.is_file(), f"Edited DOCX is missing: {UPLOADED_DOCX}"
    ledger = load_ledger(LEDGER)
    source_articles = [article for article in ledger["articles"] if "Joe Biden" in article["tags"]]
    source_articles.sort(key=lambda article: int(article["num"]), reverse=True)
    assert len(source_articles) == 17, f"Expected 17 current Joe Biden articles, found {len(source_articles)}"

    document = Document(UPLOADED_DOCX)
    assert len(document.tables) == 1, "Edited DOCX must contain exactly one table"
    table = document.tables[0]
    assert [cell.text.strip() for cell in table.rows[0].cells] == EXPECTED_HEADERS, "Unexpected DOCX headers"
    rows = table.rows[1:]
    assert len(rows) == len(source_articles), f"Expected 17 article rows, found {len(rows)}"

    records: list[dict] = []
    for source, row in zip(source_articles, rows):
        num = row.cells[0].text.strip()
        headline = row.cells[1].text.strip()
        assert num == str(source["num"]), f"NUM mismatch: expected {source['num']}, found {num}"
        assert headline == source["headline"].strip(), f"NUM {num}: immutable headline changed"
        after_tags = parse_tags(row.cells[2].text, num)
        unknown = [tag for tag in after_tags if tag not in ledger["tagMetadata"]]
        assert not unknown, f"NUM {num}: returned DOCX uses unapproved tag(s): {unknown}"
        records.append(
            {
                "num": int(source["num"]),
                "xPostUrl": source["xPostUrl"],
                "headline": source["headline"],
                "sourceUrl": source["sourceUrl"],
                "imageUrl": source["imageUrl"],
                "beforeTags": list(source["tags"]),
                "afterTags": after_tags,
            }
        )

    changed = [record for record in records if record["beforeTags"] != record["afterTags"]]
    assert len(changed) == 12, f"Expected 12 tag-set corrections, found {len(changed)}"
    biden_admin = ledger["tagMetadata"].get("Biden Administration")
    assert biden_admin == {"type": "topic", "keywords": []}, "Biden Administration metadata must remain unchanged"

    ARCHIVED_DOCX.parent.mkdir(parents=True, exist_ok=True)
    ARCHIVED_DOCX.write_bytes(UPLOADED_DOCX.read_bytes())
    docx_hash = hashlib.sha256(ARCHIVED_DOCX.read_bytes()).hexdigest()
    plan = {
        "correctionDate": "2026-10-07",
        "status": "user-approved-from-ctcrazies_Editiedjoe_biden_tag_article_list_2026-10-07.docx",
        "reviewDocument": ARCHIVED_DOCX.name,
        "reviewDocumentSha256": docx_hash,
        "expectedReviewedArticleCount": len(records),
        "expectedChangedArticleCount": len(changed),
        "records": records,
    }
    PLAN.write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    lines = [
        "# October 7, 2026 Joe Biden Historical Tag Corrections",
        "",
        "The returned DOCX was treated as the user-authoritative exact tag-set correction plan for the 17 articles that held the Joe Biden tag at review time.",
        "",
        f"- Review document: `{ARCHIVED_DOCX.name}`",
        f"- SHA-256: `{docx_hash}`",
        f"- Reviewed article rows: {len(records)}",
        f"- Tag-set corrections: {len(changed)}",
        "- Exact NUM sequence: passed",
        "- Exact immutable headlines: passed",
        "- Duplicate-tag check: passed",
        "- Final tag-count check (2–4): passed",
        "- All final tags existed in the canonical typed taxonomy: passed",
        "- No new tag declaration, rename, type change, keyword change, headline change, source URL change, X-post URL change, image change, NUM change, or ordering change was authorized.",
        "- Existing Biden Administration topic metadata was retained unchanged.",
        "",
        "## Exact corrected NUMs",
        "",
    ]
    for record in changed:
        lines.append(
            f"- **NUM {record['num']}**: `{', '.join(record['beforeTags'])}` → `{', '.join(record['afterTags'])}`"
        )
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return plan


def apply() -> None:
    assert PLAN.is_file(), f"Validated plan is missing: {PLAN}; run validate first"
    plan = json.loads(PLAN.read_text(encoding="utf-8"))
    records = plan.get("records")
    assert isinstance(records, list) and len(records) == 17, "Plan must contain exactly 17 reviewed records"

    ledger = load_ledger(LEDGER)
    original = copy.deepcopy(ledger)
    by_num = {int(article["num"]): article for article in ledger["articles"]}
    for record in records:
        article = by_num.get(int(record["num"]))
        assert article is not None, f"NUM {record['num']}: canonical article is missing"
        assert article["xPostUrl"] == record["xPostUrl"], f"NUM {record['num']}: X-post URL changed"
        assert article["headline"] == record["headline"], f"NUM {record['num']}: headline changed"
        assert article["sourceUrl"] == record["sourceUrl"], f"NUM {record['num']}: source URL changed"
        assert article["imageUrl"] == record["imageUrl"], f"NUM {record['num']}: image URL changed"
        assert article["tags"] == record["beforeTags"], f"NUM {record['num']}: current tags differ from reviewed baseline"
        unknown = [tag for tag in record["afterTags"] if tag not in ledger["tagMetadata"]]
        assert not unknown, f"NUM {record['num']}: plan has unapproved tag(s): {unknown}"
        article["tags"] = list(record["afterTags"])

    reviewed_nums = {int(record["num"]) for record in records}
    changed_nums = {
        int(after["num"])
        for before, after in zip(original["articles"], ledger["articles"])
        if before["tags"] != after["tags"]
    }
    expected_changed_nums = {
        int(record["num"])
        for record in records
        if record["beforeTags"] != record["afterTags"]
    }
    assert changed_nums == expected_changed_nums, f"Unexpected changed article scope: {changed_nums}"
    assert changed_nums.issubset(reviewed_nums), "An unreviewed article tag set would be changed"

    write_project_from_ledger(ledger, ledger["articles"][0].get("batchDate") or "2026-10-07")
    write_ledger(ledger, LEDGER)

    updated = load_ledger(LEDGER)
    updated_by_num = {int(article["num"]): article for article in updated["articles"]}
    for record in records:
        assert updated_by_num[int(record["num"])]["tags"] == record["afterTags"], f"NUM {record['num']}: final tags mismatch"
    assert sum("Joe Biden" in article["tags"] for article in updated["articles"]) == 8, "Expected 8 Joe Biden articles after correction"
    assert sum("Biden Administration" in article["tags"] for article in updated["articles"]) == 29, "Expected 29 Biden Administration articles after correction"
    print(
        json.dumps(
            {
                "reviewedArticles": len(records),
                "changedArticles": len(changed_nums),
                "changedNums": sorted(changed_nums, reverse=True),
                "joeBidenArticleCountAfter": 8,
                "bidenAdministrationArticleCountAfter": 29,
            },
            indent=2,
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=["validate", "apply"])
    args = parser.parse_args()
    if args.action == "validate":
        plan = validate()
        print(f"PASS: validated {len(plan['records'])} reviewed articles and wrote {PLAN}")
    else:
        apply()
        print("PASS: exact historical tag corrections applied and derived site files regenerated")


if __name__ == "__main__":
    main()
